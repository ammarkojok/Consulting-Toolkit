#!/usr/bin/env python3
"""Regenerate assets/library-tetp.pptx (committed, Kearney-design prototypes in
TETP colors) and refresh assets/library.json metadata.

Restricted prototypes (library.pptx, taken from real TETP decks) are indexed
here too but the file itself is gitignored; their entries are preserved as-is.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

from tetpdeck.builder import TEMPLATE  # noqa: E402
from tetpdeck.prototypes import PROTOTYPES  # noqa: E402

SKILL = Path(__file__).resolve().parent.parent
OUT = SKILL / "assets" / "library-tetp.pptx"
INDEX = SKILL / "assets" / "library.json"

# Restricted prototypes living in the gitignored library.pptx.
RESTRICTED = {
    "pathway_timeline": (1, "Two-tier rank/pathway timeline with duration chips"),
    "course_scope": (2, "Dashed-scope grouping of course options"),
    "working_group": (3, "Working-group charter matrix: purpose, members, output"),
    "course_progression": (4, "Numbered objective rows progressing across two courses"),
    "delivery_tracker": (5, "Workstream status rows with dots and progress notes"),
}


def main() -> None:
    prs = Presentation(TEMPLATE)
    layouts = {l.name: l for m in prs.slide_masters for l in m.slide_layouts}
    for sld in list(prs.slides._sldIdLst):
        prs.part.drop_rel(sld.get(qn("r:id")))
        prs.slides._sldIdLst.remove(sld)

    index: dict = {}
    for i, (name, (builder, use)) in enumerate(PROTOTYPES.items(), start=1):
        builder(prs, layouts)
        index[name] = {"file": "library-tetp.pptx", "slide": i, "use": use}

    for name, (num, use) in RESTRICTED.items():
        index[name] = {"file": "library.pptx", "slide": num, "use": use}

    prs.save(OUT)
    INDEX.write_text(json.dumps(index, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {OUT} ({len(PROTOTYPES)} prototypes) and {INDEX}")


if __name__ == "__main__":
    main()
