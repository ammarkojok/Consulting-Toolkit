#!/usr/bin/env python3
"""Generate annotated slot maps for every library prototype.

For each prototype, every text-bearing shape is filled with a short marker
([1], [2], ...) and the slide is rendered to PNG. The PNG plus a legend
(marker -> shape name -> original text) is the visual key an author reads to
address slots when filling a prototype.

Output goes to references/slotmaps/ (gitignored — the imagery and original
text derive from Restricted Access decks). Re-run whenever the library
changes:

    python3 scripts/build_slot_maps.py
"""

from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

from tetpdeck.builder import TEMPLATE  # noqa: E402
from tetpdeck.library import apply_library_slides, lib_index, list_slots  # noqa: E402

SKILL = Path(__file__).resolve().parent.parent
OUT_DIR = SKILL / "references" / "slotmaps"


def empty_deck(path: Path) -> None:
    prs = Presentation(TEMPLATE)
    for sld in list(prs.slides._sldIdLst):
        prs.part.drop_rel(sld.get(qn("r:id")))
        prs.slides._sldIdLst.remove(sld)
    prs.save(path)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    legend_md = ["# Prototype slot maps", "",
                 "Marker -> shape name -> original text. Fill slots by shape",
                 "name in `type: library` slides. Duplicate shape names are",
                 "flagged; only the first match is addressable.", ""]
    work = Path(tempfile.mkdtemp(prefix="slotmaps-"))

    for name, entry in sorted(lib_index().items()):
        try:
            slots = list_slots(name)
        except FileNotFoundError as e:
            print(f"skip {name}: {e}")
            continue
        seen: dict[str, int] = {}
        fill: dict[str, str] = {}
        legend_md += [f"## {name}", "",
                      f"`{entry['use']}`" if isinstance(entry, dict) and entry.get("use")
                      else "", "",
                      f"![{name}]({name}.png)", "",
                      "| marker | shape name | original text |",
                      "|---|---|---|"]
        for i, (shape, text) in enumerate(slots, start=1):
            dup = ""
            if shape in seen:
                dup = " (DUPLICATE name — not separately addressable)"
            else:
                seen[shape] = i
                fill[shape] = f"[{i}]"
            orig = text.replace("|", "/")[:60]
            legend_md.append(f"| [{i}] | `{shape}`{dup} | {orig} |")
        legend_md.append("")

        deck = work / f"{name}.pptx"
        empty_deck(deck)
        apply_library_slides(deck, [{"position": 0, "prototype": name,
                                     "slots": fill}])
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf",
                        "--outdir", str(work), str(deck)],
                       check=True, capture_output=True)
        subprocess.run(["pdftoppm", "-png", "-r", "80", "-singlefile",
                        str(work / f"{name}.pdf"), str(OUT_DIR / name)],
                       check=True)
        print(f"mapped {name} ({len(fill)} slots)")

    (OUT_DIR / "SLOTMAPS.md").write_text("\n".join(legend_md), encoding="utf-8")
    print(f"wrote {OUT_DIR}/SLOTMAPS.md")


if __name__ == "__main__":
    main()
