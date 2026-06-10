"""Native PowerPoint charts styled from tokens, plus the think-cell .ppttc
handoff adapter (ADR-0002)."""

from __future__ import annotations

import json
from pathlib import Path

from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

from . import tokens as T
from .geometry import Rect

KINDS = {
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "hbar": XL_CHART_TYPE.BAR_CLUSTERED,
    "stacked_bar": XL_CHART_TYPE.COLUMN_STACKED,
    "stacked_bar_100": XL_CHART_TYPE.COLUMN_STACKED_100,
    "line": XL_CHART_TYPE.LINE,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "scatter": XL_CHART_TYPE.XY_SCATTER,
}


def add_chart(slide, rect: Rect, spec: dict):
    kind = spec.get("kind", "bar")
    if kind == "scatter":
        data = XyChartData()
        for s in spec["series"]:
            sd = data.add_series(s["name"])
            for x, y in s["values"]:
                sd.add_data_point(x, y)
    else:
        data = CategoryChartData()
        data.categories = spec["categories"]
        for s in spec["series"]:
            data.add_series(s["name"], s["values"])

    gframe = slide.shapes.add_chart(
        KINDS[kind], Inches(rect.x), Inches(rect.y), Inches(rect.w), Inches(rect.h), data
    )
    chart = gframe.chart
    chart.has_title = False  # the slide title carries the conclusion
    _style(chart, spec, kind)
    return gframe


def _style(chart, spec: dict, kind: str) -> None:
    chart.font.name = T.FONT
    chart.font.size = T.SMALL_SIZE
    chart.font.color.rgb = RGBColor.from_string(T.INK)

    multi = len(spec.get("series", [])) > 1 or kind in ("pie", "doughnut")
    chart.has_legend = multi
    if multi:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = T.SMALL_SIZE

    for i, series in enumerate(chart.series):
        color = spec["series"][i].get("color") if i < len(spec.get("series", [])) else None
        color = color or T.SERIES[i % len(T.SERIES)]
        if kind in ("pie", "doughnut"):
            for j, pt in enumerate(series.points):
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = RGBColor.from_string(
                    T.SERIES[j % len(T.SERIES)]
                )
        elif kind == "line":
            series.format.line.color.rgb = RGBColor.from_string(color)
            series.format.line.width = Pt(2)
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = RGBColor.from_string(color)
            series.format.line.fill.background()

    # No gridlines; quiet axes.
    for axis_name in ("category_axis", "value_axis"):
        try:
            axis = getattr(chart, axis_name)
        except ValueError:
            continue
        axis.has_major_gridlines = False
        axis.has_minor_gridlines = False
        axis.format.line.color.rgb = RGBColor.from_string(T.GRAY)
        axis.tick_labels.font.size = T.SMALL_SIZE
        axis.tick_labels.font.name = T.FONT

    if spec.get("data_labels") and kind not in ("scatter",):
        plot = chart.plots[0]
        plot.has_data_labels = True
        plot.data_labels.font.size = T.SMALL_SIZE
        plot.data_labels.number_format_is_linked = True


# --- think-cell handoff (.ppttc) ---------------------------------------------

def emit_ppttc(out_dir: Path, deck_name: str, charts: list[dict]) -> Path | None:
    """Write a companion .ppttc for charts flagged thinkcell: true.

    The analyst opens this file on a machine with think-cell + PowerPoint and a
    template whose think-cell frames are named after each chart's `name` field.
    This artifact is a handoff, never validated by the pipeline (ADR-0002).
    """
    flagged = [c for c in charts if c.get("thinkcell")]
    if not flagged:
        return None
    entries = []
    for c in flagged:
        table = [[None] + [s["name"] for s in c["series"]]]
        for i, cat in enumerate(c.get("categories", [])):
            row = [{"string": str(cat)}]
            for s in c["series"]:
                row.append({"number": s["values"][i]})
            table.append(row)
        entries.append({"name": c.get("name", "Chart1"), "table": table})
    doc = [{"template": "REPLACE-WITH-THINKCELL-TEMPLATE.pptx", "data": entries}]
    path = out_dir / f"{deck_name}.ppttc"
    path.write_text(json.dumps(doc, indent=2, ensure_ascii=False), encoding="utf-8")
    return path
