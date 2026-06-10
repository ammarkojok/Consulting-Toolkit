"""Low-level drawing helpers. The only module that touches python-pptx shapes
directly for pattern content; everything routes through tokens + geometry."""

from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from . import tokens as T
from .geometry import Rect


def _align(rtl: bool, default=PP_ALIGN.LEFT):
    if default == PP_ALIGN.CENTER:
        return default
    return PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT


def box(slide, rect: Rect, fill: str | None = None, line: str | None = None,
        shape=MSO_SHAPE.RECTANGLE):
    sh = slide.shapes.add_shape(
        shape, Inches(rect.x), Inches(rect.y), Inches(rect.w), Inches(rect.h)
    )
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = RGBColor.from_string(line)
        sh.line.width = Pt(0.75)
    sh.text_frame.word_wrap = True
    return sh


def text(slide, rect: Rect, runs, *, size=T.BODY_SIZE, color=T.INK, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, rtl: bool = False,
         spacing: float = T.LINE_SPACING):
    """Add a borderless text box. `runs` is a string or list of paragraph dicts:
    {text, size?, color?, bold?, bullet?, level?, space_before?}."""
    tb = slide.shapes.add_textbox(
        Inches(rect.x), Inches(rect.y), Inches(rect.w), Inches(rect.h)
    )
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [{"text": runs}]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = para.get("align", _align(rtl, align))
        p.line_spacing = spacing
        if para.get("space_before"):
            p.space_before = Pt(para["space_before"])
        if para.get("level") is not None:
            p.level = para["level"]
        txt = para["text"]
        if para.get("bullet"):
            txt = ("– " if para.get("level") else "• ") + txt
        r = p.add_run()
        r.text = txt
        f = r.font
        f.name = T.FONT
        f.size = para.get("size", size)
        f.bold = para.get("bold", bold)
        f.color.rgb = RGBColor.from_string(para.get("color", color))
        if rtl:
            _set_rtl(p)
    return tb


def _set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("rtl", "1")


def accent_bar(slide, rect: Rect, color: str = T.COPPER):
    return box(slide, Rect(rect.x, rect.y, rect.w, T.ACCENT_BAR), fill=color)


def gradient_box(slide, rect: Rect, scheme: str = "navy",
                 shape=MSO_SHAPE.RECTANGLE):
    """Brand gradient block (header bars, rails) — recipe from tokens.GRADIENTS."""
    c1, c2 = T.GRADIENTS[scheme]
    sh = slide.shapes.add_shape(
        shape, Inches(rect.x), Inches(rect.y), Inches(rect.w), Inches(rect.h)
    )
    sh.shadow.inherit = False
    sh.line.fill.background()
    f = sh.fill
    f.gradient()
    f.gradient_stops[0].color.rgb = RGBColor.from_string(c1)
    f.gradient_stops[1].color.rgb = RGBColor.from_string(c2)
    f.gradient_angle = T.GRADIENT_ANGLE
    sh.text_frame.word_wrap = True
    return sh


def header_bar(slide, rect: Rect, label: str, scheme: str = "navy",
               rtl: bool = False, size=T.HEADING_SIZE):
    """Gradient header bar with centered white bold label (reference style)."""
    bar = Rect(rect.x, rect.y, rect.w, T.HEADER_BAR)
    gradient_box(slide, bar, scheme)
    text(slide, Rect(bar.x + T.PAD, bar.y, bar.w - 2 * T.PAD, bar.h),
         [{"text": label, "size": size, "bold": True, "color": T.WHITE,
           "align": PP_ALIGN.CENTER}],
         anchor=MSO_ANCHOR.MIDDLE, rtl=rtl)
    return bar


def cell(slide, rect: Rect, fill: str = T.WHITE, border: str | None = T.BORDER):
    """White content cell with a hairline border — the reference's basic unit."""
    return box(slide, rect, fill=fill, line=border)


def chip(slide, rect: Rect, label: str, fill: str = T.NAVY, color: str = T.WHITE,
         rtl: bool = False):
    """Small dark pill with white text (durations, tags)."""
    sh = box(slide, rect, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    try:
        sh.adjustments[0] = 0.5
    except Exception:
        pass
    text(slide, rect, [{"text": label, "size": T.SMALL_SIZE, "bold": True,
                        "color": color, "align": PP_ALIGN.CENTER}],
         anchor=MSO_ANCHOR.MIDDLE, rtl=rtl)
    return sh


def status_dot(slide, center_x: float, center_y: float, status: str, d: float = 0.17):
    color = T.STATUS[status]
    return box(slide, Rect(center_x - d / 2, center_y - d / 2, d, d),
               fill=color, shape=MSO_SHAPE.OVAL)


def status_legend(slide, rect: Rect, statuses: list[str], rtl: bool = False):
    seg_w = rect.w / max(1, len(statuses))
    for i, st in enumerate(statuses):
        x = rect.x + i * seg_w
        status_dot(slide, x + 0.08, rect.y + rect.h / 2, st, d=0.12)
        text(slide, Rect(x + 0.18, rect.y, seg_w - 0.2, rect.h),
             [{"text": T.STATUS_LABELS[st], "size": T.SMALL_SIZE, "color": T.MUTED}],
             anchor=MSO_ANCHOR.MIDDLE, rtl=rtl)


def bullets_to_paras(items, *, size=T.BODY_SIZE, color=T.INK):
    """Normalize spec bullet items (strings or {text, level}) to paragraph dicts."""
    out = []
    for it in items:
        if isinstance(it, str):
            it = {"text": it}
        out.append({
            "text": it["text"],
            "level": it.get("level", 0),
            "bullet": True,
            "size": size,
            "color": color,
            "bold": it.get("bold", False),
            "space_before": 4,
        })
    return out
