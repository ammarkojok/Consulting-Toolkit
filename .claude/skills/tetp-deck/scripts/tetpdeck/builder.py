"""Deck builder: validated YAML Deck Spec -> .pptx on the canonical template.

The spec never contains coordinates, fonts, or colors (beyond named accents);
this module maps slide types to layouts + patterns and hard-fails on anything
that would ship a fuckup: unknown layouts, overflow, missing fields.
"""

from __future__ import annotations

from pathlib import Path

import yaml

from . import tokens as T
from .draw import bullets_to_paras, text
from .geometry import (SPLIT_TEXT_ZONE, Rect, canvas, est_text_height)
from .patterns import PATTERNS
from .charts import emit_ppttc

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE.parent.parent / "assets" / "TETP-canonical.pptx"

# slide type -> (default layout name EN, default canvas)
PATTERN_LAYOUTS = {
    "cards": ("TETP EN Vertical", "vertical"),
    "comparison": ("TETP EN Vertical", "vertical"),
    "kpis": ("TETP EN Vertical", "vertical"),
    "process": ("TETP EN Horizontal", "horizontal"),
    "timeline": ("TETP EN Horizontal", "horizontal"),
    "chart": ("TETP EN Vertical", "vertical"),
    "table": ("TETP EN Horizontal", "horizontal"),
    "bullets": ("TETP EN Horizontal", "horizontal"),
    "tracker": ("TETP EN Horizontal", "horizontal"),
}
CANVAS_LAYOUTS = {  # canvas override -> layout that provides it
    "vertical": "TETP EN Vertical",
    "horizontal": "TETP EN Horizontal",
    "wide": "EN Wide",
    "split": "EN Split",
}
AR_EQUIV = {
    "TETP EN Cover": "TETP AR Cover",
    "TETP EN End Page": "TETP AR End Page",
    "TETP EN Vertical": "TETP AR Vertical",
    "TETP EN Horizontal": "TETP AR Horizontal",
    "TETP EN Agenda": "TETP AR Agenda",
}

# Character budgets (lint). Conservative; measured against canvas geometry.
TITLE_BUDGET = {"TETP EN Vertical": 150, "TETP AR Vertical": 150,
                "TETP EN Horizontal": 120, "TETP AR Horizontal": 120,
                "EN Split": 110, "Title Only": 150, "EN Wide": 90}


class SpecError(Exception):
    pass


class DeckBuilder:
    def __init__(self, spec_path: Path):
        self.spec_path = Path(spec_path)
        self.spec = yaml.safe_load(self.spec_path.read_text(encoding="utf-8"))
        self.errors: list[str] = []
        self.charts_for_ppttc: list[dict] = []
        meta = self.spec.get("deck", {})
        self.lang = meta.get("language", "en")
        self.rtl = self.lang == "ar"

    # --- public ---------------------------------------------------------------
    def build(self, out_path: Path) -> Path:
        from pptx import Presentation

        self._validate_top()
        if self.errors:
            raise SpecError("\n".join(self.errors))

        prs = Presentation(TEMPLATE)
        self.layouts = {l.name: l for m in prs.slide_masters for l in m.slide_layouts}
        # Drop the template's example slides (parts and rels, or the package
        # saves orphaned duplicates); generated decks start empty.
        from pptx.oxml.ns import qn
        for sld in list(prs.slides._sldIdLst):
            prs.part.drop_rel(sld.get(qn("r:id")))
            prs.slides._sldIdLst.remove(sld)

        self.lib_items = []
        position = 0
        for i, slide_spec in enumerate(self.spec["slides"], start=1):
            try:
                if slide_spec.get("type") == "library":
                    if self.rtl:
                        raise SpecError("library prototypes are LTR-only for now")
                    self.lib_items.append({
                        "position": position,
                        "prototype": slide_spec["prototype"],
                        "slots": slide_spec.get("slots") or {},
                    })
                else:
                    self._build_slide(prs, slide_spec, i)
                position += 1
            except SpecError as e:
                self.errors.append(f"slide {i}: {e}")
            except KeyError as e:
                self.errors.append(f"slide {i}: missing required field {e}")

        if self.errors:
            raise SpecError("\n".join(self.errors))

        out_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(out_path)
        if self.lib_items:
            from .library import apply_library_slides
            apply_library_slides(out_path, self.lib_items)
        ppttc = emit_ppttc(out_path.parent, out_path.stem, self.charts_for_ppttc)
        if ppttc:
            print(f"think-cell handoff written: {ppttc}")
        return out_path

    # --- internals --------------------------------------------------------------
    def _validate_top(self):
        if not isinstance(self.spec, dict) or "slides" not in self.spec:
            self.errors.append("spec must be a mapping with a 'slides' list")
            return
        if self.lang not in ("en", "ar"):
            self.errors.append(f"deck.language must be en|ar, got {self.lang!r}")

    def _layout_for(self, stype: str, slide_spec: dict) -> tuple[str, str | None]:
        if stype == "cover":
            return ("TETP EN Cover", None)
        if stype == "agenda":
            return ("TETP EN Agenda", None)
        if stype == "end":
            return ("TETP EN End Page", None)
        if stype == "two_columns":
            return ("EN Two Columns", None)
        if stype in PATTERN_LAYOUTS:
            layout, cvs = PATTERN_LAYOUTS[stype]
            if "canvas" in slide_spec:
                cvs = slide_spec["canvas"]
                if cvs not in CANVAS_LAYOUTS:
                    raise SpecError(f"unknown canvas {cvs!r}")
                layout = CANVAS_LAYOUTS[cvs]
            return (layout, cvs)
        raise SpecError(f"unknown slide type {stype!r}")

    def _resolve_language(self, layout_name: str) -> str:
        if not self.rtl:
            return layout_name
        if layout_name in AR_EQUIV:
            return AR_EQUIV[layout_name]
        raise SpecError(
            f"layout {layout_name!r} has no Arabic equivalent in the canonical "
            f"template; use canvas vertical|horizontal or types cover/agenda/end"
        )

    def _build_slide(self, prs, slide_spec: dict, idx: int):
        stype = slide_spec.get("type")
        if not stype:
            raise SpecError("slide has no 'type'")
        layout_name, cvs = self._layout_for(stype, slide_spec)
        layout_name = self._resolve_language(layout_name)
        layout = self.layouts.get(layout_name)
        if layout is None:
            raise SpecError(f"layout {layout_name!r} not found in canonical template")
        slide = prs.slides.add_slide(layout)

        fills = self._placeholder_fills(stype, slide_spec, layout_name)
        used = set()
        for ph in list(slide.placeholders):
            key = ph.placeholder_format.idx
            if key in fills and fills[key] is not None:
                ph.text = str(fills[key])
                self._restyle_placeholder(ph)
                used.add(key)
            else:
                ph._element.getparent().remove(ph._element)

        self._lint_title(stype, slide_spec, layout_name, idx)

        if stype in PATTERNS:
            cv = canvas(cvs, self.rtl) if cvs else None
            PATTERNS[stype](slide, slide_spec, cv, self.rtl)
            self._lint_pattern(stype, slide_spec, cv, idx)

        if slide_spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = slide_spec["notes"]

        if stype == "chart":
            c = dict(slide_spec["chart"])
            c.setdefault("name", f"Chart{idx}")
            self.charts_for_ppttc.append(c)

    def _placeholder_fills(self, stype: str, s: dict, layout_name: str) -> dict:
        """Map spec fields to placeholder idx per layout family."""
        if stype == "cover":
            return {0: s["title"], 1: s.get("subtitle"),
                    12: s.get("report_type"), 15: s.get("date")}
        if stype == "agenda":
            return {10: s["heading"], 22: s.get("source")}
        if stype == "end":
            return {}
        if stype == "two_columns":
            left, right = s["left"], s["right"]
            if self.rtl:
                left, right = right, left
            return {0: s["title"], 16: left.get("heading"), 17: right.get("heading"),
                    19: _as_text(left.get("body")), 20: _as_text(right.get("body")),
                    21: s.get("subtitle"), 22: s.get("source")}
        # pattern slides on canvas layouts
        fills = {0: s["title"]}
        if layout_name.endswith("Vertical"):
            fills[20] = s.get("subtitle")
            fills[22] = s.get("source")
        elif layout_name.endswith("Horizontal"):
            fills[23] = s.get("subtitle")
            fills[22] = s.get("source")
            fills[10] = None  # body placeholder unused; patterns own the canvas
        elif layout_name == "EN Split":
            fills[22] = s.get("source")
        elif layout_name == "EN Wide":
            fills[20] = s.get("subtitle")
        return fills

    def _restyle_placeholder(self, ph) -> None:
        for p in ph.text_frame.paragraphs:
            for r in p.runs:
                r.font.name = T.FONT
                if self.rtl:
                    pPr = p._p.get_or_add_pPr()
                    pPr.set("rtl", "1")
                    from pptx.enum.text import PP_ALIGN
                    p.alignment = PP_ALIGN.RIGHT

    # --- lint -------------------------------------------------------------------
    def _lint_title(self, stype, s, layout_name, idx):
        title = s.get("title", "")
        budget = TITLE_BUDGET.get(layout_name)
        if budget and len(title) > budget:
            self.errors.append(
                f"slide {idx}: title is {len(title)} chars; budget for "
                f"{layout_name} is {budget}. Shorten the action title."
            )

    def _lint_pattern(self, stype, s, cv, idx):
        if cv is None:
            return
        if stype == "cards":
            n = len(s["cards"])
            if not 2 <= n <= 6:
                self.errors.append(f"slide {idx}: cards supports 2-6 cards, got {n}")
            cols = n if n <= 4 else 3
            rows = 1 if n <= 4 else 2
            cell_w = (cv.w - T.GAP * (cols - 1)) / cols - 2 * T.PAD
            cell_h = (cv.h - T.GAP * (rows - 1)) / rows - 2 * T.PAD
            for c in s["cards"]:
                body = c.get("body", "")
                body = " ".join(body) if isinstance(body, list) else body
                h = (est_text_height(c["heading"], cell_w, T.HEADING_SIZE.pt)
                     + est_text_height(body, cell_w, T.BODY_SIZE.pt) + 0.2)
                if h > cell_h:
                    self.errors.append(
                        f"slide {idx}: card '{c['heading'][:30]}' overflows its cell "
                        f"(~{h:.1f}in needed, {cell_h:.1f}in available). Trim copy or "
                        f"use canvas: horizontal."
                    )
        if stype == "kpis":
            n = len(s["kpis"])
            if not 2 <= n <= 5:
                self.errors.append(f"slide {idx}: kpis supports 2-5 items, got {n}")
        if stype == "process" and not 2 <= len(s["steps"]) <= 6:
            self.errors.append(f"slide {idx}: process supports 2-6 steps")
        if stype == "timeline" and not 2 <= len(s["milestones"]) <= 8:
            self.errors.append(f"slide {idx}: timeline supports 2-8 milestones")
        if stype == "tracker":
            n = len(s["rows"])
            if not 2 <= n <= 7:
                self.errors.append(f"slide {idx}: tracker supports 2-7 rows, got {n}")
            bad = [r.get("status") for r in s["rows"]
                   if r.get("status", "not_started") not in T.STATUS]
            if bad:
                self.errors.append(
                    f"slide {idx}: unknown tracker status {bad[0]!r}; use "
                    f"{', '.join(T.STATUS)}"
                )
        if stype == "comparison":
            for side in ("left", "right"):
                blts = s[side].get("bullets", [])
                total = sum(len(b if isinstance(b, str) else b["text"]) for b in blts)
                if total > 700:
                    self.errors.append(
                        f"slide {idx}: {side} panel has ~{total} chars of bullets; "
                        f"keep under 700."
                    )


def _as_text(body) -> str | None:
    if body is None:
        return None
    if isinstance(body, list):
        return "\n".join(b if isinstance(b, str) else b["text"] for b in body)
    return str(body)
