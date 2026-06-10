"""Core 8 pattern renderers. Each pattern draws content onto a layout canvas;
geometry comes from geometry.py, styling from tokens.py — patterns decide
arrangement only. Every renderer takes (slide, spec, canvas_rect, rtl)."""

from __future__ import annotations

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from . import tokens as T
from .draw import accent_bar, box, bullets_to_paras, text
from .geometry import Rect, columns, est_text_height, grid, hsplit, vsplit
from .charts import add_chart


def cards(slide, spec, cv: Rect, rtl: bool):
    items = spec["cards"]
    max_cols = len(items) if len(items) <= 4 else 3
    cells = grid(cv, len(items), max_cols=max_cols, gap=T.GAP, rtl=rtl)
    for cell, item in zip(cells, items):
        box(slide, cell, fill=T.WHITE)
        accent_bar(slide, cell, item.get("accent", T.COPPER))
        inner = cell.inset(T.PAD)
        paras = [{"text": item["heading"], "size": T.HEADING_SIZE, "bold": True,
                  "color": T.NAVY}]
        if item.get("body"):
            body = item["body"] if isinstance(item["body"], list) else [item["body"]]
            for b in body:
                paras.append({"text": b if isinstance(b, str) else b["text"],
                              "size": T.BODY_SIZE, "space_before": 5})
        text(slide, Rect(inner.x, inner.y + 0.10, inner.w, inner.h - 0.10),
             paras, rtl=rtl)


def comparison(slide, spec, cv: Rect, rtl: bool):
    panels = [spec["left"], spec["right"]]
    colors = [T.NAVY, T.COPPER]
    for rect, panel, color in zip(columns(cv, 2, T.GAP, rtl), panels, colors):
        head_h = 0.46
        box(slide, Rect(rect.x, rect.y, rect.w, head_h), fill=color)
        text(slide, Rect(rect.x + T.PAD, rect.y, rect.w - 2 * T.PAD, head_h),
             [{"text": panel["heading"], "size": T.HEADING_SIZE, "bold": True,
               "color": T.WHITE}], anchor=MSO_ANCHOR.MIDDLE, rtl=rtl)
        body = Rect(rect.x, rect.y + head_h, rect.w, rect.h - head_h)
        box(slide, body, fill=T.WHITE)
        text(slide, body.inset(T.PAD),
             bullets_to_paras(panel.get("bullets", [])), rtl=rtl)


def kpis(slide, spec, cv: Rect, rtl: bool):
    items = spec["kpis"]
    band_h = min(1.9, cv.h)
    band = Rect(cv.x, cv.y + (cv.h - band_h) / 2 if spec.get("centered", True) else cv.y,
                cv.w, band_h)
    col_w = (band.w - T.GAP * (len(items) - 1)) / len(items)
    for rect, item in zip(columns(band, len(items), T.GAP, rtl), items):
        box(slide, rect, fill=T.WHITE)
        accent_bar(slide, rect, item.get("accent", T.NAVY))
        # Step the value size down when it would wrap the column.
        value = str(item["value"])
        vsize = T.KPI_VALUE_SIZE
        from pptx.util import Pt
        while len(value) * 0.0095 * vsize.pt > col_w - 2 * T.PAD and vsize.pt > 20:
            vsize = Pt(vsize.pt - 4)
        paras = [{"text": value, "size": vsize, "bold": True,
                  "color": T.NAVY, "align": PP_ALIGN.CENTER}]
        if item.get("delta"):
            paras.append({"text": item["delta"], "size": T.SMALL_SIZE,
                          "color": T.MUTED, "align": PP_ALIGN.CENTER})
        paras.append({"text": item["label"], "size": T.KPI_LABEL_SIZE,
                      "color": T.INK, "align": PP_ALIGN.CENTER, "space_before": 4})
        text(slide, rect.inset(T.PAD), paras, anchor=MSO_ANCHOR.MIDDLE, rtl=rtl)


def process(slide, spec, cv: Rect, rtl: bool):
    from pptx.enum.shapes import MSO_SHAPE
    steps = spec["steps"]
    chev_h = 0.55
    body_h = cv.h - chev_h - 0.12
    cols = columns(Rect(cv.x, cv.y, cv.w, cv.h), len(steps), T.GAP, rtl)
    for i, (rect, step) in enumerate(zip(cols, steps)):
        shape = MSO_SHAPE.PENTAGON if i == 0 and not rtl else MSO_SHAPE.CHEVRON
        chev = box(slide, Rect(rect.x, rect.y, rect.w, chev_h),
                   fill=T.SERIES[i % len(T.SERIES)], shape=shape)
        tf = chev.text_frame
        tf.word_wrap = True
        text(slide, Rect(rect.x + 0.12, rect.y, rect.w - 0.24, chev_h),
             [{"text": step["heading"], "size": T.BODY_SIZE, "bold": True,
               "color": T.WHITE, "align": PP_ALIGN.CENTER}],
             anchor=MSO_ANCHOR.MIDDLE, rtl=rtl)
        if step.get("body"):
            body = step["body"] if isinstance(step["body"], list) else [step["body"]]
            text(slide, Rect(rect.x, rect.y + chev_h + 0.12, rect.w, body_h),
                 bullets_to_paras(body), rtl=rtl)


def timeline(slide, spec, cv: Rect, rtl: bool):
    miles = spec["milestones"]
    line_y = cv.y + cv.h * 0.42
    box(slide, Rect(cv.x + 0.1, line_y, cv.w - 0.2, 0.025), fill=T.GRAY)
    cols = columns(Rect(cv.x, cv.y, cv.w, cv.h), len(miles), 0.08, rtl)
    dot = 0.16
    for i, (rect, m) in enumerate(zip(cols, miles)):
        cx = rect.x + rect.w / 2
        from pptx.enum.shapes import MSO_SHAPE
        done = m.get("done", False)
        box(slide, Rect(cx - dot / 2, line_y - dot / 2 + 0.0125, dot, dot),
            fill=T.NAVY if done else T.WHITE,
            line=None if done else T.NAVY, shape=MSO_SHAPE.OVAL)
        text(slide, Rect(rect.x, line_y - 0.62, rect.w, 0.5),
             [{"text": m["date"], "size": T.SMALL_SIZE, "bold": True,
               "color": T.COPPER, "align": PP_ALIGN.CENTER}],
             anchor=MSO_ANCHOR.BOTTOM, rtl=rtl)
        text(slide, Rect(rect.x + 0.05, line_y + 0.22, rect.w - 0.1, cv.h - (line_y - cv.y) - 0.3),
             [{"text": m["label"], "size": T.BODY_SIZE, "align": PP_ALIGN.CENTER}],
             rtl=rtl)


def chart_takeaways(slide, spec, cv: Rect, rtl: bool):
    if spec.get("takeaways"):
        chart_rect, side = hsplit(cv, [0.66, 0.34], T.GAP, rtl)
        box(slide, side, fill=T.WHITE)
        accent_bar(slide, side, T.NAVY)
        inner = side.inset(T.PAD)
        paras = [{"text": spec.get("takeaways_heading", "Key takeaways"),
                  "size": T.HEADING_SIZE, "bold": True, "color": T.NAVY}]
        paras += bullets_to_paras(spec["takeaways"])
        text(slide, Rect(inner.x, inner.y + 0.1, inner.w, inner.h - 0.1), paras, rtl=rtl)
    else:
        chart_rect = cv
    pad = box(slide, chart_rect, fill=T.WHITE)
    add_chart(slide, chart_rect.inset(0.18, 0.16), spec["chart"])


def table(slide, spec, cv: Rect, rtl: bool):
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    cols_spec = spec["columns"]
    rows_spec = spec["rows"]
    if rtl:
        cols_spec = list(reversed(cols_spec))
        rows_spec = [list(reversed(r)) for r in rows_spec]
    nrows, ncols = len(rows_spec) + 1, len(cols_spec)
    h = min(cv.h, 0.38 * nrows)
    gf = slide.shapes.add_table(nrows, ncols, Inches(cv.x), Inches(cv.y),
                                Inches(cv.w), Inches(h))
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    for j, name in enumerate(cols_spec):
        cell = tbl.cell(0, j)
        cell.text = str(name)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(T.NAVY)
        _style_cell(cell, bold=True, color=T.WHITE, rtl=rtl)
    for i, row in enumerate(rows_spec, start=1):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            # YAML parses bare Yes/No as booleans; render them back as words.
            if val is True:
                val = "Yes"
            elif val is False:
                val = "No"
            cell.text = "" if val is None else str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(
                T.WHITE if i % 2 else "F7F7F7"
            )
            _style_cell(cell, bold=False, color=T.INK, rtl=rtl)


def _style_cell(cell, *, bold: bool, color: str, rtl: bool):
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    cell.margin_left = cell.margin_right = Pt(6)
    cell.margin_top = cell.margin_bottom = Pt(3)
    for p in cell.text_frame.paragraphs:
        p.alignment = PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT
        for r in p.runs:
            r.font.name = T.FONT
            r.font.size = T.BODY_SIZE
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(color)


def bullets_pattern(slide, spec, cv: Rect, rtl: bool):
    """Bullet hierarchy drawn on a canvas (used when a native body placeholder
    isn't the target). Optionally splits into two balanced columns."""
    items = spec["bullets"]
    paras = bullets_to_paras(items)
    if spec.get("two_column") and len(items) > 5:
        half = (len(paras) + 1) // 2
        for rect, chunk in zip(columns(cv, 2, T.SECTION_GAP, rtl),
                               (paras[:half], paras[half:])):
            text(slide, rect, chunk, rtl=rtl)
    else:
        text(slide, cv, paras, rtl=rtl)


PATTERNS = {
    "cards": cards,
    "comparison": comparison,
    "kpis": kpis,
    "process": process,
    "timeline": timeline,
    "chart": chart_takeaways,
    "table": table,
    "bullets": bullets_pattern,
}
