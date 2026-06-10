#!/usr/bin/env python3
"""Extract manifest.json from the canonical template.

The Manifest is the template's interface: the authoring LLM reads this file,
never the .pptx binary. Re-run after any change to the canonical template.
"""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
TEMPLATE = SKILL / "assets" / "TETP-canonical.pptx"
OUT = SKILL / "assets" / "manifest.json"

EMU_IN = 914400


def rect(shape) -> dict | None:
    try:
        if shape.left is None:
            return None
        return {
            "x": round(shape.left / EMU_IN, 2),
            "y": round(shape.top / EMU_IN, 2),
            "w": round(shape.width / EMU_IN, 2),
            "h": round(shape.height / EMU_IN, 2),
        }
    except Exception:
        return None


def char_budget(r: dict | None, font_pt: float) -> int | None:
    """Rough capacity: chars per line * lines that fit."""
    if not r:
        return None
    cpl = int(r["w"] / (0.0075 * font_pt))
    lines = int(r["h"] / (font_pt / 72 * 1.3))
    return max(cpl, cpl * lines)


def main() -> None:
    prs = Presentation(TEMPLATE)
    manifest = {
        "template": TEMPLATE.name,
        "slide_size_in": [round(prs.slide_width / EMU_IN, 2),
                          round(prs.slide_height / EMU_IN, 2)],
        "theme": {
            "fonts": {"major": "Arial", "minor": "Arial"},
            "colors": {
                "ink": "1E1E1E", "white": "FFFFFF", "navy": "3A3B68",
                "copper": "C4996C", "gray": "D2D2D2", "tan": "C5A783",
                "mauve": "B6AFB0", "indigo": "66689B", "deep": "505282",
            },
        },
        "masters": [],
    }
    for m in prs.slide_masters:
        lang = "ar" if any("AR" in l.name for l in m.slide_layouts) else "en"
        master = {"language": lang, "layouts": []}
        for lay in m.slide_layouts:
            entry = {"name": lay.name, "placeholders": []}
            for ph in lay.placeholders:
                r = rect(ph)
                prompt = ph.text_frame.text.strip()[:80] if ph.has_text_frame else ""
                font_pt = 18 if ph.placeholder_format.idx == 0 else 11
                entry["placeholders"].append({
                    "idx": ph.placeholder_format.idx,
                    "type": str(ph.placeholder_format.type).split(" ")[0],
                    "rect_in": r,
                    "char_budget": char_budget(r, font_pt),
                    "prompt": prompt,
                })
            master["layouts"].append(entry)
        manifest["masters"].append(master)

    OUT.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
